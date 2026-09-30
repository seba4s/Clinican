"""Lectura de las fichas de peluquería que CLINICAN llenaba en Excel [A-9].

Cada archivo es una ficha: una mascota y una visita (hoja con «Fecha», «Mascota»,
«Raza», «Sexo», «Propietario», «C.C.», «Cel/ Tel», «Dirección», la descripción del
corte con las casillas marcadas con X, el valor y una nota).

Los datos se buscan por su rótulo y no por la celda, para tolerar filas o columnas
corridas entre un archivo y otro. Aquí solo se lee; la validación y el guardado están
en ``clinican.servicios.importacion.importar_fichas``.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from clinican.dominio.formato import sin_tildes

# Rótulo normalizado -> dato de la ficha
_DATOS = {
    "fecha": "fecha",
    "mascota": "mascota",
    "nombre mascota": "mascota",
    "raza": "raza",
    "sexo": "sexo",
    "propietario": "propietario",
    "cc": "cedula",
    "cedula": "cedula",
    "cel/tel": "celular",
    "celular": "celular",
    "telefono": "celular",
    "tel": "celular",
    "direccion": "direccion",
}

# Casillas de la descripción del corte (rótulo normalizado -> clave)
MARCAS = {
    "tijera": "tijera",
    "despunte": "despunte",
    "bano cepillado": "bano_cepillado",
    "bano medicado": "bano_medicado",
    "bajar 1 cm": "bajar_1cm",
    "bajar 1/2 cm": "bajar_medio_cm",
    "patas rasuradas": "patas_rasuradas",
    "barbas": "barbas",
    "bigotes": "bigotes",
    "copete": "copete",
    "orejas": "orejas",
    "cola": "cola",
    "bano antipulgas": "bano_antipulgas",
    "desparasitacion": "desparasitacion",
}
NOMBRE_MARCA = {
    "tijera": "Tijera", "despunte": "Despunte", "bano_cepillado": "Baño - cepillado", "bano_medicado": "Baño medicado",
    "bajar_1cm": "Bajar 1 cm", "bajar_medio_cm": "Bajar 1/2 cm", "patas_rasuradas": "Patas rasuradas",
    "barbas": "Barbas", "bigotes": "Bigotes", "copete": "Copete", "orejas": "Orejas", "cola": "Cola",
    "bano_antipulgas": "Baño antipulgas", "desparasitacion": "Desparasitación",
}
_FIN_NOTA = ("nuestro servicio incluye", "clausulas", "he leido")


@dataclass
class FichaExcel:
    archivo: str
    fecha: str | None = None
    mascota: str | None = None
    raza: str | None = None
    sexo: str | None = None
    propietario: str | None = None
    cedula: str | None = None
    celulares: list[str] = field(default_factory=list)
    direccion: str | None = None
    marcas: dict[str, str] = field(default_factory=dict)  # clave -> lo escrito en la casilla ("X", "LEON")
    valor: int | None = None
    nota: str | None = None


def rotulo(texto) -> str:
    """"Cel/ Tel" -> "cel/tel", "C.C." -> "cc", "Baño - cepillado" -> "bano cepillado"."""
    t = sin_tildes(str(texto or "")).replace(".", "").replace(":", " ")
    t = re.sub(r"\s*/\s*", "/", t)
    t = re.sub(r"[^a-z0-9/ ]", " ", t)
    return " ".join(t.split())


def _texto(valor) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = " ".join(str(valor).split())
    return texto or None


def valor_en_pesos(texto: str | None) -> int | None:
    """"VALOR 45 MIL" -> 45000, "$45.000" -> 45000, "45000" -> 45000."""
    if not texto:
        return None
    t = sin_tildes(texto)
    m = re.search(r"(\d[\d.,]*)\s*(mil|k)?\b", t.split("valor", 1)[-1])
    if not m:
        return None
    numero = int(re.sub(r"[.,]", "", m.group(1)))
    return numero * 1000 if m.group(2) else numero


def _telefonos(texto: str | None) -> list[str]:
    if not texto:
        return []
    partes = re.split(r"[/;,]|\sy\s|\s-\s", texto)
    numeros = [re.sub(r"\D", "", p) for p in partes]
    return [n for n in numeros if 7 <= len(n) <= 10]


def _sexo(texto: str | None) -> str | None:
    t = sin_tildes(texto or "").strip()
    if t.startswith(("h", "f")):
        return "HEMBRA"
    if t.startswith("m"):
        return "MACHO"
    return None


def leer(ruta: str | Path) -> FichaExcel:
    """Lee una ficha. Lanza ``ValueError`` si el archivo no se puede abrir o no parece una ficha."""
    from openpyxl import load_workbook

    ruta = Path(ruta)
    try:
        libro = load_workbook(ruta, data_only=True)
    except Exception as e:  # dañado, abierto en otro programa o de otro formato
        raise ValueError(f"No se pudo abrir el archivo: {e}") from None
    try:
        hoja = libro.active
        celdas: dict[tuple[int, int], object] = {}
        for fila in hoja.iter_rows():
            for c in fila:
                if _texto(c.value) is not None:
                    celdas[(c.row, c.column)] = c.value
    finally:
        libro.close()

    ficha = FichaExcel(archivo=ruta.name)
    rotulos = {pos: rotulo(v) for pos, v in celdas.items() if isinstance(v, str)}

    def a_la_derecha(fila: int, col: int, hasta: int = 3):
        for c in range(col + 1, col + 1 + hasta):
            v = celdas.get((fila, c))
            if v is not None and rotulos.get((fila, c)) not in _DATOS and rotulos.get((fila, c)) not in MARCAS:
                return v
        return None

    encontrados = 0
    for (fila, col), r in sorted(rotulos.items()):
        if r in _DATOS:
            dato = _DATOS[r]
            valor = a_la_derecha(fila, col)
            encontrados += 1
            if dato == "celular":
                ficha.celulares += _telefonos(_texto(valor))
                abajo = celdas.get((fila + 1, col + 1))  # segundo número en la fila de abajo
                if abajo is not None and (fila + 1, col) not in celdas:
                    ficha.celulares += _telefonos(_texto(abajo))
            elif dato == "fecha":
                ficha.fecha = _texto(valor)
            elif dato == "sexo":
                ficha.sexo = _sexo(_texto(valor))
            elif getattr(ficha, dato) is None:
                setattr(ficha, dato, _texto(valor))
        elif r in MARCAS:
            valor = _texto(celdas.get((fila, col + 1)))
            if valor:
                ficha.marcas[MARCAS[r]] = valor
        elif r.startswith(("descripcion del corte", "valor")):
            # «DESCRIPCION DEL CORTE: VALOR 45 MIL», «VALOR 45 MIL» o «Valor» con el monto al lado
            ficha.valor = (ficha.valor or valor_en_pesos(str(celdas[(fila, col)]))
                           or valor_en_pesos(_texto(a_la_derecha(fila, col))))
        elif r == "nota":
            partes = [_texto(a_la_derecha(fila, col))]
            for siguiente in range(fila + 1, fila + 4):
                textos = [(c, rotulos.get((siguiente, c), "")) for (f, c) in celdas if f == siguiente]
                if not textos or any(t.startswith(_FIN_NOTA) or t in _DATOS or t in MARCAS for _c, t in textos):
                    break
                partes += [_texto(celdas[(siguiente, c)]) for c, _t in sorted(textos)]
            ficha.nota = " ".join(p for p in partes if p) or None
    if encontrados < 3:
        raise ValueError("No parece una ficha de peluquería de CLINICAN (no se encontraron «Fecha», «Mascota», "
                         "«Propietario»…).")
    ficha.celulares = list(dict.fromkeys(ficha.celulares))  # sin repetidos, en orden
    return ficha
