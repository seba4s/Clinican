"""Lectura de la plantilla de importación y generación de la plantilla vacía (sección 11).

Aquí solo se lee el archivo; la validación y el guardado están en
``clinican.servicios.importacion``. Cuando se consiga un Excel real [A-9],
basta con adaptar ``leer`` a su formato.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from clinican.dominio.formato import sin_tildes

HOJA = "Mascotas"

# (clave interna, encabezado de la plantilla, obligatoria, ancho)
COLUMNAS: list[tuple[str, str, bool, int]] = [
    ("cedula", "Cédula del propietario *", True, 20),
    ("propietario", "Nombre del propietario *", True, 28),
    ("celular1", "Celular 1 *", True, 15),
    ("celular2", "Celular 2", False, 15),
    ("direccion", "Dirección *", True, 30),
    ("mascota", "Nombre de la mascota *", True, 22),
    ("raza", "Raza *", True, 30),
    ("tamano", "Tamaño (Pequeña, Mediana o Grande)", False, 22),
    ("pelaje", "Pelaje complicado (Sí o No)", False, 16),
    ("edad_anios", "Edad (años)", False, 11),
    ("edad_meses", "Edad (meses)", False, 11),
    ("ultima_visita", "Fecha última visita (AAAA-MM-DD)", False, 20),
    ("observaciones", "Observaciones", False, 40),
]

INSTRUCCIONES = [
    "CÓMO LLENAR LA PLANTILLA DE IMPORTACIÓN — CLINICAN",
    None,
    "• Una fila por mascota. Si un propietario tiene varias mascotas, repita su cédula en cada fila.",
    "• Las columnas con * son obligatorias.",
    "• Raza: elija de la lista. Si la raza no existe, primero agréguela en Configuración › Razas.",
    "• Tamaño: obligatorio para perros mestizos. En las demás razas, déjelo vacío para usar el de la raza.",
    "• Pelaje complicado: obligatorio para perros mestizos (Sí o No).",
    "• Fecha de última visita: formato AAAA-MM-DD (por ejemplo 2026-03-26) o como fecha de Excel.",
    "• Si la cédula ya existe en el sistema, la mascota se agrega a ese propietario (sus datos no se cambian).",
    "• Los consentimientos NO se importan: cada propietario debe aceptarlos antes de su primer turno.",
    None,
    "Ejemplo de fila:",
    "1098765432 | María Pérez | 3001234567 | | Calle 10 # 5-20 | Toby | Shih Tzu | | | 3 | 0 | 2026-02-15 | "
    "Nervioso con el secador",
]


def _clave_encabezado(texto) -> str:
    """"Tamaño (Pequeña, ...)" -> "tamano". Ignora tildes, mayúsculas, asteriscos y paréntesis."""
    limpio = sin_tildes(str(texto or "")).replace("*", "")
    limpio = re.sub(r"\(.*?\)", "", limpio)
    return " ".join(limpio.split())


_POR_ENCABEZADO = {_clave_encabezado(titulo): clave for clave, titulo, _o, _a in COLUMNAS}


def _texto(valor) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    return texto or None


def leer(ruta: str | Path) -> list[tuple[int, dict]]:
    """Devuelve [(número de fila en Excel, {clave: valor en texto o None})], sin las filas vacías.

    Lanza ``ValueError`` con un mensaje claro si el archivo no se puede leer o le faltan columnas.
    """
    from openpyxl import load_workbook

    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except Exception as e:  # archivo dañado, abierto en otro programa o de otro formato
        raise ValueError(f"No se pudo abrir el archivo de Excel: {e}") from None
    try:
        hoja = libro[HOJA] if HOJA in libro.sheetnames else libro.active
        filas = hoja.iter_rows(values_only=True)
        encabezados = next(filas, None) or ()
        posiciones: dict[str, int] = {}
        for i, titulo in enumerate(encabezados):
            clave = _POR_ENCABEZADO.get(_clave_encabezado(titulo))
            if clave and clave not in posiciones:
                posiciones[clave] = i
        faltan = [titulo for clave, titulo, obligatoria, _a in COLUMNAS if obligatoria and clave not in posiciones]
        if faltan:
            raise ValueError(
                "Al archivo le faltan las columnas: " + ", ".join(t.replace(" *", "") for t in faltan)
                + ". Use la plantilla de importación de CLINICAN."
            )
        resultado = []
        for numero, fila in enumerate(filas, start=2):
            datos = {clave: _texto(fila[i]) if i < len(fila) else None for clave, i in posiciones.items()}
            for clave, *_ in COLUMNAS:
                datos.setdefault(clave, None)
            if any(v is not None for v in datos.values()):
                resultado.append((numero, datos))
        return resultado
    finally:
        libro.close()


def generar_plantilla(ruta: str | Path, razas: list[str]) -> Path:
    """Crea la plantilla vacía con la lista de razas activas para elegir."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    ruta = Path(ruta)
    libro = Workbook()
    hoja = libro.active
    hoja.title = HOJA
    verde = PatternFill("solid", fgColor="B7CF53")
    for j, (_clave, titulo, _o, ancho) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=j, value=titulo)
        celda.font = Font(bold=True)
        celda.fill = verde
        hoja.column_dimensions[celda.column_letter].width = ancho
    hoja.freeze_panes = "A2"

    hoja_razas = libro.create_sheet("Razas")
    for i, nombre in enumerate(razas, start=1):
        hoja_razas.cell(row=i, column=1, value=nombre)

    ultima = 2000
    validaciones = [
        (f"=Razas!$A$1:$A${max(len(razas), 1)}", f"G2:G{ultima}"),
        ('"Pequeña,Mediana,Grande"', f"H2:H{ultima}"),
        ('"Sí,No"', f"I2:I{ultima}"),
    ]
    for formula, rango in validaciones:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.add(rango)
        hoja.add_data_validation(dv)

    instrucciones = libro.create_sheet("Instrucciones")
    for i, linea in enumerate(INSTRUCCIONES, start=1):
        instrucciones.cell(row=i, column=1, value=linea)
    instrucciones["A1"].font = Font(bold=True, size=14)
    instrucciones.column_dimensions["A"].width = 110

    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)
    return ruta
