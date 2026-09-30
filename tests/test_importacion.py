"""Sección 11: plantilla de importación, validación y agrupación por cédula."""

import pytest
from openpyxl import load_workbook

from clinican.dominio.errores import DatoInvalido, PermisoDenegado
from clinican.servicios import importacion
from clinican.servicios import propietarios as sp


def _llenar(ruta, filas):
    libro = load_workbook(ruta)
    hoja = libro["Mascotas"]
    for i, fila in enumerate(filas, start=2):
        for j, valor in enumerate(fila, start=1):
            hoja.cell(row=i, column=j, value=valor)
    libro.save(ruta)


FILAS = [
    # cédula, dueño, cel1, cel2, dirección, mascota, raza, tamaño, pelaje, años, meses, última visita, obs
    ("1098765432", "María Pérez", "3012345678", None, "Calle 10", "Toby", "Shih Tzu", None, None, 3, 0, "2026-02-15", None),
    ("1.098.765.432", "María Pérez", "3012345678", None, "Calle 10", "Luna", "shih tzu", None, None, 1, 6, None, None),
    ("5555555", "Pedro Gómez", "3150000000", None, "Cra 2", "Max", "Perro mestizo (sin raza)", "Mediana", "No", None, None, None, "Nervioso"),
    ("6666666", "Sin Tamaño", "3160000000", None, "Cra 3", "Rex", "Perro mestizo (sin raza)", None, "No", None, None, None, None),
    ("7777777", "Raza Rara", "3170000000", None, "Cra 4", "Kiki", "Dálmata", None, None, None, None, None, None),
    ("8888888", "Mal Celular", "12", None, "Cra 5", "Bobo", "Husky", None, None, None, None, None, None),
    ("1098765432", "María Pérez", "3012345678", None, "Calle 10", "toby", "Shih Tzu", None, None, None, None, None, None),
]


@pytest.fixture
def plantilla(conn, tmp_path):
    ruta = importacion.generar_plantilla(conn, tmp_path / "plantilla_importacion.xlsx")
    _llenar(ruta, FILAS)
    return ruta


def test_plantilla_tiene_columnas_y_razas(conn, tmp_path):
    ruta = importacion.generar_plantilla(conn, tmp_path / "p.xlsx")
    libro = load_workbook(ruta)
    assert {"Mascotas", "Razas", "Instrucciones"} <= set(libro.sheetnames)
    assert libro["Mascotas"]["A1"].value.startswith("Cédula del propietario")
    assert libro["Razas"].max_row == 14


def test_validar_no_guarda(conn, empleada, plantilla):
    r = importacion.importar(conn, empleada, plantilla, solo_validar=True)
    assert r.filas_leidas == 7 and r.mascotas_nuevas == 3 and r.propietarios_nuevos == 2
    assert conn.execute("SELECT COUNT(*) FROM mascotas").fetchone()[0] == 0


def test_importar_agrupa_por_cedula_y_reporta_rechazos(conn, empleada, plantilla):
    r = importacion.importar(conn, empleada, plantilla)
    assert r.importado
    assert r.mascotas_nuevas == 3 and r.propietarios_nuevos == 2
    rechazadas = dict(r.rechazadas)
    assert set(rechazadas) == {5, 6, 7, 8}
    assert "tamaño" in rechazadas[5]
    assert "Dálmata" in rechazadas[6]
    assert "celular" in rechazadas[7].lower()
    assert "Toby" in rechazadas[8] or "toby" in rechazadas[8]
    maria = sp.buscar(conn, empleada, "1098765432")
    assert len(maria) == 1
    assert sorted(m["nombre"] for m in sp.mascotas(conn, empleada, maria[0]["id"])) == ["Luna", "Toby"]
    toby = [m for m in sp.mascotas(conn, empleada, maria[0]["id"]) if m["nombre"] == "Toby"][0]
    assert toby["fecha_ultima_visita"] == "2026-02-15"
    assert "IMPORTAR_EXCEL" in [f[0] for f in conn.execute("SELECT accion FROM auditoria")]


def test_propietario_existente_recibe_mascotas(conn, empleada, plantilla):
    pid = sp.crear(conn, empleada, "Pedro Gómez", "5555555", "3150000000", None, "Cra 2")
    r = importacion.importar(conn, empleada, plantilla)
    assert r.propietarios_existentes == 1 and r.propietarios_nuevos == 1
    assert [m["nombre"] for m in sp.mascotas(conn, empleada, pid)] == ["Max"]


def test_importar_dos_veces_no_duplica(conn, empleada, plantilla):
    importacion.importar(conn, empleada, plantilla)
    r = importacion.importar(conn, empleada, plantilla)
    assert r.mascotas_nuevas == 0
    assert conn.execute("SELECT COUNT(*) FROM mascotas").fetchone()[0] == 3


def test_archivo_sin_columnas(conn, empleada, tmp_path):
    from openpyxl import Workbook

    ruta = tmp_path / "malo.xlsx"
    libro = Workbook()
    libro.active.append(["Nombre", "Otra cosa"])
    libro.save(ruta)
    with pytest.raises(DatoInvalido, match="faltan"):
        importacion.importar(conn, empleada, ruta)


def test_importar_requiere_sesion(conn, plantilla):
    with pytest.raises(PermisoDenegado):
        importacion.importar(conn, None, plantilla)
